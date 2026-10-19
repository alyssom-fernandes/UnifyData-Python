"""Testes das regras do fechamento (nucleo.py). Rodam com `python -m unittest`, sem dependências além do app."""
import io
import sys
import unittest
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))
import nucleo as N  # noqa: E402

EXEMPLOS = RAIZ / 'exemplos'


class Arquivo(io.BytesIO):
    """Imita o arquivo enviado pela tela (name e getvalue)."""
    def __init__(self, nome, conteudo):
        super().__init__(conteudo)
        self.name = nome


def do_disco(caminho):
    return Arquivo(caminho.name, caminho.read_bytes())


def rh_exemplo():
    return [do_disco(p) for p in sorted(EXEMPLOS.glob('Empregados *.xlsx'))]


def linx_exemplo():
    return [do_disco(EXEMPLOS / 'LOJA.csv'), do_disco(EXEMPLOS / 'RESTAURANTE.csv')]


class Leitura(unittest.TestCase):
    def test_planilhas_do_rh_de_exemplo(self):
        hr, avisos = N.parse_hr_files(rh_exemplo())
        self.assertEqual(avisos, [])
        self.assertEqual(len(hr), 20)
        self.assertEqual(sorted(hr['Empresa'].unique()),
                         ['Auto Peças Horizonte', 'Hotel Primavera - Filial', 'Hotel Primavera - Matriz', 'Posto Serra Azul'])
        self.assertTrue((hr['CPF'].str.len() == 11).all())

    def test_relatorios_do_linx_de_exemplo(self):
        linx, avisos, periodos, _ = N.parse_linx_files(linx_exemplo())
        self.assertEqual(avisos, [])
        self.assertEqual(len(linx), 34)
        self.assertEqual(periodos, ['01/09/2026 a 30/09/2026'])
        primeiro = linx.iloc[0].to_dict()
        self.assertEqual(primeiro, {'Nome': 'ANA CAROLINA MENDES', 'Vencido': 11.1, 'Total': 18.5, 'Origem': 'LOJA', 'Linha': N.SEM_LINHA})

    def test_linha_de_total_sem_vencido_e_com_milhar(self):
        texto = ('Conta: 1 Responsável: 1 - ANA;;\r\nTotal;;18,50;;\r\n'
                 'Conta: 1 Responsável: 2 - BRUNO;;\r\nTotal;1.234,56;2.000,10;;\r\n'
                 'Total geral;1.234,56;2.018,60;;\r\n').encode('cp1252')
        linx, _, _, _ = N.parse_linx_files([Arquivo('LOJA.csv', texto)])
        self.assertEqual(linx[['Nome', 'Vencido', 'Total']].values.tolist(), [['ANA', 0.0, 18.5], ['BRUNO', 1234.56, 2000.1]])

    def test_total_geral_do_relatorio_e_lido_por_origem(self):
        _, _, _, declarados = N.parse_linx_files(linx_exemplo())
        self.assertEqual(declarados, {'LOJA': 1385.6, 'RESTAURANTE': 1296.2})

    def test_total_geral_que_bate_com_a_soma_dos_responsaveis(self):
        texto = ('Conta: 1 Responsável: 1 - ANA;;\r\nTotal;;18,50;;\r\n'
                 'Conta: 1 Responsável: 2 - BRUNO;;\r\nTotal;1.234,56;2.000,10;;\r\n'
                 'Total geral;1.234,56;2.018,60;;\r\n').encode('cp1252')
        linx, _, _, declarados = N.parse_linx_files([Arquivo('LOJA.csv', texto)])
        c = N.conferencia(N.totais_por_loja(linx), [], declarados)
        self.assertEqual(declarados, {'LOJA': 2018.6})
        self.assertEqual((c['consumo'], c['no_relatorio'], c['divergentes']), (2018.6, 2018.6, []))

    def test_total_geral_que_nao_bate_vira_divergencia(self):
        # O bloco do BRUNO não traz a linha "Total": o parser não o reconhece e o valor some da soma
        texto = ('Conta: 1 Responsável: 1 - ANA;;\r\nTotal;;18,50;;\r\n'
                 'Conta: 1 Responsável: 2 - BRUNO;;\r\nSubtotal;1.234,56;2.000,10;;\r\n'
                 'Total geral;1.234,56;2.018,60;;\r\n').encode('cp1252')
        linx, _, _, declarados = N.parse_linx_files([Arquivo('LOJA.csv', texto)])
        c = N.conferencia(N.totais_por_loja(linx), [], declarados)
        self.assertEqual(c['consumo'], 18.5)
        self.assertEqual(c['no_relatorio'], 2018.6)
        self.assertEqual(c['divergentes'], [('LOJA', 18.5, 2018.6)])

    def test_relatorio_sem_total_geral_nao_e_conferido(self):
        texto = 'Conta: 1 Responsável: 1 - ANA;;\r\nTotal;;18,50;;\r\n'.encode('cp1252')
        linx, _, _, declarados = N.parse_linx_files([Arquivo('LOJA.csv', texto)])
        c = N.conferencia(N.totais_por_loja(linx), [], declarados)
        self.assertEqual(declarados, {})
        self.assertIsNone(c['no_relatorio'])
        self.assertEqual(c['divergentes'], [])

    def test_csv_do_rh_com_titulo_aspas_e_rodape(self):
        texto = ('RELAÇÃO DE FUNCIONÁRIOS\nNome;CPF;Endereço\n"SILVA, JOÃO";123.456.789-01;"Rua A, 1"\n'
                 'TOTAL: 1;;\n').encode('utf-8')
        hr, avisos = N.parse_hr_files([Arquivo('Empregados Loja.csv', texto)])
        self.assertEqual(avisos, [])
        self.assertEqual(hr.to_dict('records'), [{'Nome': 'SILVA, JOÃO', 'CPF': '12345678901', 'Empresa': 'Loja'}])

    def test_planilha_ilegivel_fica_de_fora_com_aviso_claro(self):
        with self.assertLogs('nucleo', level='WARNING') as registro:
            hr, avisos = N.parse_hr_files([Arquivo('Empregados Loja.xlsx', b'isto nao e uma planilha')])
        self.assertTrue(hr.empty)
        self.assertEqual(avisos, ['**Empregados Loja.xlsx**: não parece uma planilha .xls, .xlsx ou .csv válida, o arquivo ficou de fora.'])
        self.assertIn('Empregados Loja.xlsx', registro.output[0], 'o detalhe técnico vai para o console, não para a tela')


class Cruzamento(unittest.TestCase):
    def test_exemplos_um_ligado_por_acento_e_cinco_divergencias(self):
        hr, _ = N.parse_hr_files(rh_exemplo())
        linx, _, _, _ = N.parse_linx_files(linx_exemplo())
        divergencias, ligados = N.encontrar_divergencias(hr, linx)
        self.assertEqual(ligados, 1)
        self.assertIn('DÉBORA LIMA FONSECA', set(linx['Nome']), 'DEBORA (Linx) vira DÉBORA (RH)')
        tipos = [d['Tipo'] for d in divergencias]
        self.assertEqual(tipos, ['Nome Semelhante'] + ['Não Encontrado'] * 4)
        self.assertEqual(divergencias[0]['HR_Nome'], 'RAFAEL CUNHA BATISTA')

    def test_conferencia_total_do_linx_menos_ignorado_e_o_total_da_planilha(self):
        # O fluxo da demonstração: aprova o nome parecido e ignora as quatro pessoas fora do RH
        hr, _ = N.parse_hr_files(rh_exemplo())
        linx, _, periodos, declarados = N.parse_linx_files(linx_exemplo())
        totais = N.totais_por_loja(linx)
        divergencias, _ = N.encontrar_divergencias(hr, linx)
        ignorados = []
        for d in divergencias:
            if d['Tipo'] == 'Nome Semelhante':
                N.renomear(linx, d['Linx_Nome'], d['HR_Nome'])
            else:
                ignorados.append({'nome': d['Linx_Nome'], 'valor': d['Valor']})
                linx = linx[~N.pelo_nome(linx, d['Linx_Nome'])]
        total = float(N.consolidar(hr, linx, list(totais))['Total Desconto'].sum())
        c = N.conferencia(totais, ignorados, declarados)
        self.assertEqual([o for o, _ in c['por_origem']], ['LOJA', 'RESTAURANTE'])
        self.assertEqual(c['no_relatorio'], c['consumo'], 'a soma lida fecha com o Total geral impresso pelo Linx')
        self.assertEqual(c['divergentes'], [])
        self.assertEqual(c['a_descontar'], round(total, 2), 'total do Linx − ignorado = total da planilha')
        self.assertEqual(N.mes_do_periodo(periodos), 'setembro de 2026')
        self.assertEqual(N.mes_do_periodo(['16/09/2026 a 15/10/2026']), '16/09/2026 a 15/10/2026')

    def test_acento_ambiguo_nao_liga_sozinho(self):
        hr = pd.DataFrame([{'Nome': 'JOSÉ SILVA', 'CPF': '1', 'Empresa': 'A'}, {'Nome': 'JOSE SILVA', 'CPF': '2', 'Empresa': 'B'}])
        linx = pd.DataFrame([{'Nome': 'JOSÈ SILVA', 'Vencido': 0.0, 'Total': 5.0, 'Origem': 'LOJA'}])
        divergencias, ligados = N.encontrar_divergencias(hr, linx)
        self.assertEqual(ligados, 0)
        self.assertEqual(divergencias[0]['Tipo'], 'Nome Semelhante')

    def test_homonimo_vira_revisao_e_o_desconto_nao_dobra(self):
        hr = pd.DataFrame([{'Nome': 'ANA', 'CPF': '11111111111', 'Empresa': 'Posto'},
                           {'Nome': 'ANA', 'CPF': '22222222222', 'Empresa': 'Hotel'}])
        linx = pd.DataFrame([{'Nome': 'ANA', 'Vencido': 0.0, 'Total': 10.0, 'Origem': 'LOJA'}])
        divergencias, _ = N.encontrar_divergencias(hr, linx)
        self.assertEqual(divergencias[0]['Tipo'], 'Repetido')
        self.assertEqual([r for _, r in divergencias[0]['Opcoes']], ['Posto · CPF 111.111.111-11', 'Hotel · CPF 222.222.222-22'])

        sem_escolha = N.consolidar(hr, linx, ['LOJA'])
        self.assertEqual(sem_escolha['Total Desconto'].sum(), 20.0, 'sem a revisão, o homônimo pagaria duas vezes')
        # Caminho real do app: a escolha na revisão prende o consumo à linha escolhida (fixar_linha)
        N.fixar_linha(linx, 'ANA', 0)
        self.assertIsNone(N.divergencia_repetido(hr, linx, 'ANA'), 'escolhido, o cartão some')
        com_escolha = N.consolidar(hr, linx, ['LOJA'])
        self.assertEqual(com_escolha['Total Desconto'].tolist(), [10.0, 0.0])

    def test_cpf_fixa_a_linha_e_a_escolha_do_homonimo_nao_leva_junto(self):
        # Maria (Linx) é, pelo CPF, o João da Alfa; o João do Linx é escolhido como o da Beta
        hr = pd.DataFrame([{'Nome': 'JOAO PEDRO SANTOS', 'CPF': '11111111111', 'Empresa': 'Alfa'},
                           {'Nome': 'JOAO PEDRO SANTOS', 'CPF': '22222222222', 'Empresa': 'Beta'}])
        linx = pd.DataFrame([{'Nome': 'JOAO PEDRO SANTOS', 'Vencido': 0.0, 'Total': 80.0, 'Origem': 'LOJA'},
                             {'Nome': 'MARIA', 'Vencido': 0.0, 'Total': 20.0, 'Origem': 'LOJA'}])
        N.fixar_linha(linx, 'MARIA', 0)
        repetido = N.divergencia_repetido(hr, linx, 'JOAO PEDRO SANTOS')
        self.assertEqual(repetido['Valor'], 80.0, 'o consumo já fixado pelo CPF não entra no cartão do homônimo')
        N.fixar_linha(linx, 'JOAO PEDRO SANTOS', 1)
        self.assertIsNone(N.divergencia_repetido(hr, linx, 'JOAO PEDRO SANTOS'))
        merged = N.consolidar(hr, linx, ['LOJA'])
        self.assertEqual(merged[['Empresa', 'Total Desconto']].values.tolist(), [['Alfa', 20.0], ['Beta', 80.0]])

    def test_consolidar_inclui_quem_nao_consumiu(self):
        hr = pd.DataFrame([{'Nome': 'ANA', 'CPF': '', 'Empresa': 'Posto'}, {'Nome': 'BRUNO', 'CPF': '', 'Empresa': 'Posto'}])
        linx = pd.DataFrame([{'Nome': 'ANA', 'Vencido': 2.0, 'Total': 3.0, 'Origem': 'LOJA'}])
        merged = N.consolidar(hr, linx, ['LOJA', 'RESTAURANTE'])
        self.assertEqual(merged[['Nome', 'Vencido__LOJA', 'Total__RESTAURANTE', 'Total Desconto']].values.tolist(),
                         [['ANA', 2.0, 0.0, 3.0], ['BRUNO', 0.0, 0.0, 0.0]])


class Exportacao(unittest.TestCase):
    def setUp(self):
        hr = pd.DataFrame([{'Nome': '=HYPERLINK("x")', 'CPF': '12345678901', 'Empresa': 'Obras: Sede'},
                           {'Nome': 'ÂNGELA', 'CPF': '', 'Empresa': 'Obras: Sede'},
                           {'Nome': 'ZÉLIA', 'CPF': '', 'Empresa': 'History'}])
        linx = pd.DataFrame([{'Nome': '=HYPERLINK("x")', 'Vencido': 1.5, 'Total': 2.0, 'Origem': 'LOJA'}])
        self.merged = N.consolidar(hr, linx, ['LOJA'])

    def test_excel_abas_validas_rodape_e_nome_como_texto(self):
        wb = load_workbook(N.gerar_excel(self.merged, ['LOJA']))
        self.assertEqual(wb.sheetnames, ['History (empresa)', 'Obras - Sede'])
        ws = wb['Obras - Sede']
        self.assertEqual([c.value for c in ws[1]], ['Funcionário', 'CPF', 'Vencido LOJA', 'Total LOJA', 'Total Desconto'])
        self.assertEqual(ws['A2'].value, '=HYPERLINK("x")')
        self.assertEqual(ws['A2'].data_type, 's')
        self.assertEqual(ws['B2'].value, '123.456.789-01')
        self.assertEqual(ws['A3'].value, 'ÂNGELA')
        self.assertTrue(ws['A3'].font.color.rgb.endswith('808080'), 'quem não consumiu fica em cinza')
        self.assertEqual([c.value for c in ws[4]], ['TOTAL GERAL', None, 1.5, 2.0, 2.0])
        self.assertEqual(ws.freeze_panes, 'A2')

    def test_excel_da_demonstracao_avisa_que_os_dados_sao_ficticios(self):
        wb = load_workbook(N.gerar_excel(self.merged, ['LOJA'], demo=True))
        ws = wb['Obras - Sede']
        self.assertEqual(ws.cell(ws.max_row, 1).value, N.AVISO_DEMO)
        self.assertEqual([c.value for c in ws[4]], ['TOTAL GERAL', None, 1.5, 2.0, 2.0], 'o aviso fica depois do total')
        self.assertEqual(ws.oddHeader.center.text, 'DEMONSTRAÇÃO: dados fictícios')
        self.assertLess(ws.column_dimensions['A'].width, 48, 'o aviso não alarga a coluna dos nomes')
        sem_demo = load_workbook(N.gerar_excel(self.merged, ['LOJA']))['Obras - Sede']
        self.assertEqual(sem_demo.max_row, 4)

    def test_nome_do_arquivo_pelo_mes_dos_vencimentos(self):
        hoje = date(2026, 10, 13)
        self.assertEqual(N.nome_do_arquivo(['01/09/2026 a 30/09/2026'], hoje=hoje), 'Fechamento_UnifyData_2026-09')
        self.assertEqual(N.nome_do_arquivo(['16/09/2026 a 15/10/2026'], hoje=hoje), 'Fechamento_UnifyData_2026-09_a_2026-10')
        self.assertEqual(N.nome_do_arquivo([], hoje=hoje), 'Fechamento_UnifyData_2026-10-13', 'sem período, vale a data de hoje')
        self.assertEqual(N.nome_do_arquivo(['01/09/2026 a 30/09/2026'], demo=True, hoje=hoje), 'Fechamento_UnifyData_2026-09_DEMO')
        _, _, periodos, _ = N.parse_linx_files(linx_exemplo())
        self.assertEqual(N.nome_do_arquivo(periodos, hoje=hoje), 'Fechamento_UnifyData_2026-09')

    def test_csv_para_o_excel_brasileiro(self):
        csv = N.gerar_csv(self.merged, ['LOJA']).decode('utf-8')
        self.assertTrue(csv.startswith(N.BOM))
        linhas = csv[1:].split('\r\n')
        self.assertEqual(linhas[0], 'Empresa;Funcionário;CPF;Vencido LOJA;Total LOJA;Total Desconto')
        self.assertEqual(linhas[1], 'History;ZÉLIA;;0,00;0,00;0,00')
        self.assertEqual(linhas[2], 'Obras: Sede;"\'=HYPERLINK(""x"")";123.456.789-01;1,50;2,00;2,00')

    def test_nome_de_aba(self):
        usados = set()
        self.assertEqual(N.nome_de_aba('Obras: Nova Sede / "Bloco A"', usados), 'Obras - Nova Sede - "Bloco A"')
        self.assertEqual(N.nome_de_aba('[Matriz]?*', usados), '(Matriz)')
        longo = 'Empresa com um nome comprido demais para uma aba'
        self.assertEqual(N.nome_de_aba(longo, usados), 'Empresa com um nome comprido')
        self.assertEqual(N.nome_de_aba(longo, usados), 'Empresa com um nome (2)')
        # Corte em palavra inteira, sem espaço nem hífen no fim
        self.assertEqual(N.nome_de_aba('Cooperativa Agroindustrial dos Produtores Rurais', usados), 'Cooperativa Agroindustrial dos')
        self.assertEqual(N.nome_de_aba('Transportes Rodoviários - Filial Norte', usados), 'Transportes Rodoviários')

    def test_textos(self):
        self.assertEqual(N.nome_exibicao('MARIA DAS DORES DE SOUZA'), 'Maria das Dores de Souza')
        self.assertEqual(N.normalizar('  José  da Silva '), 'JOSE DA SILVA')
        self.assertEqual(N.lista_por_extenso(['A', 'B', 'C']), 'A, B e C')
        self.assertEqual(N.empresa_do_arquivo('Empregados Posto Central.xls'), 'Posto Central')
        self.assertEqual(N.brl(1234.5), 'R$ 1.234,50')


if __name__ == '__main__':
    unittest.main()
