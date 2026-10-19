"""
UnifyData — nucleo.py
Regras do fechamento, sem interface: leitura das planilhas do RH e dos
relatórios do Linx, cruzamento pelo nome, consolidação e exportação.
A tela (unifydata.py) e os testes (tests/) usam só estas funções.
"""
import csv
import hashlib
import io
import logging
import math
import re
import unicodedata
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from thefuzz import fuzz, process

NOMES_COLUNA = ['nome', 'nome completo', 'funcionario', 'funcionário']
PARTICULAS = {'da', 'das', 'de', 'do', 'dos', 'e'}
SEMELHANCA_MINIMA = 85
SEM_LINHA = -1  # consumo ainda ligado pelo nome, não por uma linha do RH
BOM = chr(0xFEFF)
AVISO_DEMO = 'Demonstração com dados fictícios, gerada com os arquivos de exemplo do UnifyData. Não é um fechamento real.'
log = logging.getLogger(__name__)


# ──────────────────────────────────────────────
#  Texto
# ──────────────────────────────────────────────
def normalizar(texto):
    """Maiúsculas, sem acento e com espaços únicos — para comparar nomes."""
    sem_acento = unicodedata.normalize('NFD', str(texto)).encode('ascii', 'ignore').decode()
    return ' '.join(sem_acento.upper().split())


def brl(valor):
    return f"R$ {valor:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')


def so_digitos(texto):
    return re.sub(r'\D', '', str(texto or ''))


def formatar_cpf(cpf):
    d = so_digitos(cpf)
    return f"{d[:3]}.{d[3:6]}.{d[6:9]}-{d[9:]}" if len(d) == 11 else (cpf or '')


def nome_exibicao(nome):
    """Nome em caixa alta → "Maria da Silva" (só para exibição)."""
    partes = str(nome).lower().split()
    return ' '.join(p if i and p in PARTICULAS else re.sub(r"(^|['’\"(-])(\w)", lambda m: m.group(1) + m.group(2).upper(), p)
                    for i, p in enumerate(partes))


def lista_por_extenso(itens):
    itens = list(itens)
    return ' e '.join(itens) if len(itens) < 3 else ', '.join(itens[:-1]) + ' e ' + itens[-1]


def texto_seguro(valor):
    """Texto que começa com = + - @ viraria fórmula no Excel; o apóstrofo neutraliza (usado no CSV)."""
    texto = str(valor or '')
    return "'" + texto if texto[:1] in ('=', '+', '-', '@', '\t', '\r') else texto


def decodificar(conteudo):
    """Relatórios do Linx vêm em Windows-1252; arquivos salvos depois podem estar em UTF-8."""
    try:
        return conteudo.decode('utf-8-sig')
    except UnicodeDecodeError:
        return conteudo.decode('cp1252', errors='replace')


def id_de(texto):
    return hashlib.md5(texto.encode('utf-8')).hexdigest()[:10]


def empresa_do_arquivo(nome_arquivo):
    """"Empregados Posto Central.xls" → "Posto Central"."""
    base = nome_arquivo.rsplit('.', 1)[0]
    return base[11:].strip() if base.lower().startswith('empregados ') else base


# ──────────────────────────────────────────────
#  Planilhas do RH
# ──────────────────────────────────────────────
def find_column(df, targets):
    """Localiza uma coluna pelo nome (cabeçalho ou primeiros valores)."""
    targets_lower = [t.strip().lower() for t in targets]
    for col in df.columns:
        if str(col).strip().lower() in targets_lower:
            return col
    # Busca nos primeiros 20 valores de cada coluna
    for col in df.columns:
        for val in df[col].head(20).values:
            if pd.notna(val) and str(val).strip().lower() in targets_lower:
                return col
    return None


def linha_do_cabecalho(df, col, targets):
    """Posição da linha onde está o cabeçalho da coluna (-1 quando o cabeçalho já é o nome da coluna)."""
    alvos = {t.strip().lower() for t in targets}
    if str(col).strip().lower() in alvos:
        return -1
    for pos, val in enumerate(df[col].head(20).values):
        if pd.notna(val) and str(val).strip().lower() in alvos:
            return pos
    return -1


def extrair_funcionarios(df, empresa):
    nome_col = find_column(df, NOMES_COLUNA)
    cpf_col = find_column(df, ['cpf'])
    if nome_col is None:
        return []
    # Títulos acima do cabeçalho ("RELAÇÃO DE FUNCIONÁRIOS") não são gente
    inicio = linha_do_cabecalho(df, nome_col, NOMES_COLUNA) + 1
    registros = []
    for _, row in df.iloc[inicio:].iterrows():
        nome = ' '.join(str(row[nome_col]).strip().upper().split())
        if not nome or nome in ('NAN', 'NONE', 'NOME', 'NOME COMPLETO', 'FUNCIONARIO', 'FUNCIONÁRIO') or len(nome) < 3:
            continue
        # Títulos, totais e rodapés ("TOTAL: 12", "SISTEMA LICENCIADO PARA: ...") não são funcionários
        if ':' in nome or not re.search(r'[A-ZÀ-Ý]', nome):
            continue
        cpf = ''
        valor = row[cpf_col] if cpf_col is not None else None
        if valor is not None and pd.notna(valor):
            cpf = so_digitos(int(valor) if isinstance(valor, float) and valor.is_integer() else valor)
            if 9 <= len(cpf) < 11:
                cpf = cpf.zfill(11)  # CPF salvo como número perde os zeros à esquerda
        registros.append({'Nome': nome, 'CPF': cpf, 'Empresa': empresa})
    return registros


def ler_csv_rh(texto):
    """CSV de RH em DataFrame sem cabeçalho fixo (linhas de título não atrapalham), com ; ou , como separador."""
    primeiras = texto.splitlines()[:30]
    cabecalho = next((l for l in primeiras if 'cpf' in l.lower()), primeiras[0] if primeiras else '')
    sem_aspas = re.sub(r'"[^"]*"', '', cabecalho)
    sep = ';' if sem_aspas.count(';') >= sem_aspas.count(',') else ','
    return pd.DataFrame(list(csv.reader(io.StringIO(texto), delimiter=sep)))


def parse_hr_files(hr_files):
    """Lê arquivos de RH (XLS, XLSX ou CSV) e retorna (DataFrame com Nome, CPF e Empresa, avisos)."""
    records, avisos = [], []
    for file in hr_files:
        try:
            base = empresa_do_arquivo(file.name)
            ext = file.name.rsplit('.', 1)[-1].lower()
            antes = len(records)

            if ext == 'csv':
                records += extrair_funcionarios(ler_csv_rh(decodificar(file.getvalue())), base)
            else:
                xls = pd.ExcelFile(file)
                for sheet in xls.sheet_names:
                    empresa = f"{base} - {sheet}" if (len(xls.sheet_names) > 1 and not sheet.lower().startswith('plan')) else base
                    records += extrair_funcionarios(pd.read_excel(xls, sheet_name=sheet), empresa)

            if len(records) == antes:
                avisos.append(f"**{file.name}**: não encontrei funcionários (colunas Nome e CPF), o arquivo ficou de fora.")
        except Exception:
            # Qualquer arquivo ruim só fica de fora: a tela mostra o aviso e o detalhe técnico vai para o console
            log.warning('Planilha do RH ilegível: %s', file.name, exc_info=True)
            avisos.append(f"**{file.name}**: não parece uma planilha .xls, .xlsx ou .csv válida, o arquivo ficou de fora.")

    df = pd.DataFrame(records) if records else pd.DataFrame(columns=['Nome', 'CPF', 'Empresa'])
    return df.drop_duplicates(subset=['Nome', 'CPF', 'Empresa']).reset_index(drop=True), avisos


# ──────────────────────────────────────────────
#  Relatórios do Linx
# ──────────────────────────────────────────────
def parse_linx_files(linx_files):
    """
    Lê relatórios CSV do Linx (Pendências por Responsável).
    Retorna (DataFrame com Nome, Vencido, Total e Origem, avisos, períodos de vencimento, "Total geral" de cada
    relatório) — capturando os DOIS valores da linha Total separadamente. O "Total geral" é o valor que o próprio
    Linx imprime no fim do relatório, guardado por origem para a conferência comparar com a soma dos responsáveis.
    """
    def to_float(s):
        try:
            return float(re.sub(r'[R$\s]', '', s).replace('.', '').replace(',', '.'))
        except (ValueError, AttributeError):
            return 0.0

    records, avisos, periodos, declarados = [], [], [], {}
    for file in linx_files:
        try:
            content      = decodificar(file.getvalue())
            origem       = re.sub(r'\.(csv|txt)$', '', file.name, flags=re.IGNORECASE)
            current_nome = None
            antes        = len(records)

            # Período do relatório, para identificar o fechamento: "Vencimento: 01/09/2026 a 30/09/2026"
            m = re.search(r'Vencimento:\s*(\d{2}/\d{2}/\d{4}\s*a\s*\d{2}/\d{2}/\d{4})', content, re.IGNORECASE)
            if m and ' '.join(m.group(1).split()) not in periodos:
                periodos.append(' '.join(m.group(1).split()))

            for line in content.split('\n'):
                if 'Respons' in line:
                    m = re.search(r'Respons[^\s:]*:\s*\d+\s*-\s*([^;\r\n]+)', line, re.IGNORECASE)
                    if m:
                        current_nome = ' '.join(m.group(1).strip().upper().split())

                if re.match(r'total\s+geral', line.strip(), re.IGNORECASE):
                    # O Total geral vem depois do último responsável: é o valor do Linx para conferir a soma lida
                    campos = [p.strip() for p in line.split(';') if p.strip()]
                    if len(campos) >= 2:
                        declarados[origem] = declarados.get(origem, 0.0) + to_float(campos[2] if len(campos) >= 3 else campos[1])
                    current_nome = None
                    continue

                if current_nome and line.strip().lower().startswith('total'):
                    campos = [p.strip() for p in line.split(';')]
                    # "Total;;18,50": sem vencido. Nos demais casos, campos vazios são ignorados.
                    parts = ['Total', '0', campos[2]] if len(campos) > 2 and not campos[1] and campos[2] else [p for p in campos if p]
                    vencido = to_float(parts[1]) if len(parts) >= 2 else 0.0
                    total   = to_float(parts[2]) if len(parts) >= 3 else vencido
                    if total > 0:
                        records.append({'Nome': current_nome, 'Vencido': vencido, 'Total': total, 'Origem': origem})
                    current_nome = None

            if len(records) == antes:
                avisos.append(f"**{file.name}**: nenhum consumo encontrado, o arquivo ficou de fora.")
        except Exception:
            log.warning('Relatório do Linx ilegível: %s', file.name, exc_info=True)
            avisos.append(f"**{file.name}**: não parece um relatório do Linx em .csv, o arquivo ficou de fora.")

    df = pd.DataFrame(records) if records else pd.DataFrame(columns=['Nome', 'Vencido', 'Total', 'Origem'])
    # Linha do RH dona do consumo: SEM_LINHA enquanto o cruzamento é pelo nome; o índice da linha quando o CPF já decidiu
    df['Linha'] = SEM_LINHA
    return df, avisos, periodos, {o: round(v, 2) for o, v in declarados.items()}


# ──────────────────────────────────────────────
#  Cruzamento pelo nome
# ──────────────────────────────────────────────
def linhas_fixas(linx_df):
    """Linha do RH de cada consumo (SEM_LINHA quando ainda vale o nome); tabelas sem a coluna valem pelo nome."""
    return linx_df['Linha'] if 'Linha' in linx_df else pd.Series(SEM_LINHA, index=linx_df.index)


def pelo_nome(linx_df, nome):
    """Consumos com esse nome que ainda dependem do nome para achar a pessoa (o CPF não fixou a linha)."""
    return (linx_df['Nome'] == nome) & (linhas_fixas(linx_df) == SEM_LINHA)


def renomear(linx_df, nome_antigo, nome_novo):
    linx_df.loc[pelo_nome(linx_df, nome_antigo), 'Nome'] = nome_novo


def fixar_linha(linx_df, nome, linha):
    """O CPF apontou a pessoa: o consumo vai para essa linha do RH, seja qual for a escolha feita depois para o nome."""
    alvo = pelo_nome(linx_df, nome)
    if 'Linha' not in linx_df:
        linx_df['Linha'] = SEM_LINHA
    linx_df.loc[alvo, 'Linha'] = int(linha)


def divergencia_repetido(hr_df, linx_df, nome):
    """Nome do Linx que aparece em mais de uma linha do RH (homônimos ou a mesma pessoa em duas empresas)."""
    linhas = hr_df[hr_df['Nome'] == nome]
    if len(linhas) < 2:
        return None
    consumo = linx_df[pelo_nome(linx_df, nome)]
    if consumo.empty:
        return None
    return {
        'id': id_de('rep:' + nome), 'Tipo': 'Repetido', 'Linx_Nome': nome,
        'Opcoes': [(int(i), f"{r['Empresa']} · CPF {formatar_cpf(r['CPF']) or 'não informado'}") for i, r in linhas.iterrows()],
        'Origem': ', '.join(consumo['Origem'].unique()), 'Valor': float(consumo['Total'].sum()),
    }


def encontrar_divergencias(hr_df, linx_df):
    """
    Liga sozinho os nomes que só diferem por acento, caixa ou espaço (quando há uma única grafia
    no RH) e lista o resto para revisão. Altera linx_df e devolve (divergências, nomes ligados).
    """
    hr_names = set(hr_df['Nome'])
    por_chave = {}
    for n in hr_names:
        por_chave.setdefault(normalizar(n), set()).add(n)
    ligados = 0
    for nome in linx_df['Nome'].unique():
        alvos = por_chave.get(normalizar(nome), set())
        if nome not in hr_names and len(alvos) == 1:
            renomear(linx_df, nome, next(iter(alvos)))
            ligados += 1

    hr_lista = hr_df['Nome'].tolist()
    divergencias = []
    for nome in linx_df['Nome'].unique():
        if nome in hr_names:
            repetido = divergencia_repetido(hr_df, linx_df, nome)
            if repetido:
                divergencias.append(repetido)
            continue
        best, score = process.extractOne(nome, hr_lista, scorer=fuzz.token_sort_ratio)
        consumo = linx_df[linx_df['Nome'] == nome]
        divergencias.append({
            'id': id_de(nome), 'Linx_Nome': nome, 'HR_Nome': best, 'Score': score,
            'Tipo': 'Nome Semelhante' if score >= SEMELHANCA_MINIMA else 'Não Encontrado',
            'Origem': ', '.join(consumo['Origem'].unique()), 'Valor': float(consumo['Total'].sum()),
        })
    ordem = {'Repetido': 0, 'Nome Semelhante': 1, 'Não Encontrado': 2}
    divergencias.sort(key=lambda d: (ordem[d['Tipo']], d['Linx_Nome']))
    return divergencias, ligados


def consolidar(hr_df, linx_df, origens):
    """Uma linha por funcionário do RH, com vencido e total por origem e o Total Desconto."""
    vazio = linx_df is None or linx_df.empty
    por_nome = pd.DataFrame() if vazio else linx_df[linhas_fixas(linx_df) == SEM_LINHA]
    por_linha = pd.DataFrame() if vazio else linx_df[linhas_fixas(linx_df) != SEM_LINHA]

    def make_pivot(df, chave, col):
        if df.empty:
            return pd.DataFrame(columns=[chave])
        p = df.pivot_table(index=chave, columns='Origem', values=col, aggfunc='sum', fill_value=0)
        # Renomeia as colunas ANTES do reset_index para evitar conflito com a chave
        p.columns = [f'{col}__{o}' for o in p.columns]
        return p.reset_index()

    base = hr_df.assign(_linha=hr_df.index)
    merged = (base.merge(make_pivot(por_nome, 'Nome', 'Vencido'), on='Nome', how='left')
                  .merge(make_pivot(por_nome, 'Nome', 'Total'), on='Nome', how='left'))
    valores = [f'{p}__{o}' for o in origens for p in ('Vencido', 'Total')]
    for col in valores:
        if col not in merged:
            merged[col] = 0.0
    merged[valores] = merged[valores].fillna(0.0)
    # Consumo cuja pessoa já foi definida (pelo CPF ou pela escolha na revisão): soma direto na linha certa, sem passar pelo nome
    if not por_linha.empty:
        for prefixo in ('Vencido', 'Total'):
            extra = make_pivot(por_linha, 'Linha', prefixo).set_index('Linha')
            for col in extra.columns:
                merged[col] = merged[col] + merged['_linha'].map(extra[col]).fillna(0.0)
    merged['Total Desconto'] = sum(merged[f'Total__{o}'] for o in origens) if origens else 0.0
    return merged.drop(columns='_linha')


def totais_por_loja(linx_df):
    """Soma dos responsáveis lidos em cada relatório do Linx, guardada antes de qualquer divergência ser ignorada."""
    if linx_df is None or linx_df.empty:
        return {}
    return {o: round(float(v), 2) for o, v in linx_df.groupby('Origem', sort=False)['Total'].sum().items()}


def divergencias_do_total_geral(totais_linx, declarados):
    """
    Relatórios cuja soma dos responsáveis lidos difere do "Total geral" impresso pelo Linx: [(origem, lido, no relatório)].
    Relatório sem a linha "Total geral" não entra: não há com o que comparar.
    """
    return [(o, lido, declarados[o]) for o, lido in totais_linx.items()
            if o in declarados and abs(round(lido - declarados[o], 2)) >= 0.01]


def conferencia(totais_linx, ignorados, declarados=None):
    """
    A conta do fechamento como numa fita de calculadora: o total de cada loja no Linx, a soma, o que foi
    ignorado e o que vai para a folha (que precisa bater com o total da planilha). Com os "Total geral" lidos dos
    relatórios, confere também a soma lida contra o que o Linx imprimiu e lista as divergências.
    """
    declarados = declarados or {}
    consumo = round(sum(totais_linx.values()), 2)
    ignorado = round(sum(x['valor'] for x in ignorados), 2)
    completo = bool(totais_linx) and all(o in declarados for o in totais_linx)
    return {'por_origem': list(totais_linx.items()), 'consumo': consumo, 'ignorado': ignorado,
            'a_descontar': round(consumo - ignorado, 2),
            'no_relatorio': round(sum(declarados[o] for o in totais_linx), 2) if completo else None,
            'divergentes': divergencias_do_total_geral(totais_linx, declarados)}


def mes_do_periodo(periodos):
    """["01/09/2026 a 30/09/2026"] → "setembro de 2026"; vencimentos em mais de um mês ficam como vieram."""
    meses = set(re.findall(r'\d{2}/(\d{2})/(\d{4})', ' '.join(periodos or [])))
    if len(meses) != 1:
        return '; '.join(periodos or [])
    mes, ano = next(iter(meses))
    nomes = ['janeiro', 'fevereiro', 'março', 'abril', 'maio', 'junho', 'julho', 'agosto', 'setembro', 'outubro', 'novembro', 'dezembro']
    return f'{nomes[int(mes) - 1]} de {ano}'


# ──────────────────────────────────────────────
#  Exportação
# ──────────────────────────────────────────────
def cortar_nome(texto, limite):
    """Corta no limite sem deixar palavra pela metade nem espaço, hífen ou apóstrofo no fim."""
    if len(texto) <= limite:
        return texto
    corte = texto[:limite]
    espaco = corte.rfind(' ')
    if texto[limite] != ' ' and espaco >= limite / 2:
        corte = corte[:espaco]
    return re.sub(r"[\s'-]+$", '', corte)


def nome_de_aba(nome, usados):
    """Excel não aceita \\ / ? * [ ] : em nome de aba, nem mais de 31 caracteres, nomes repetidos ou "History"."""
    base = re.sub(r'\s*[:/\\]\s*', ' - ', str(nome))
    base = re.sub(r'[?*]', '', base).replace('[', '(').replace(']', ')')
    base = ' '.join(base.split()).strip("'") or 'Empresa'
    if base.lower() == 'history':
        base = 'History (empresa)'
    base = cortar_nome(base, 31)
    candidato, n = base, 2
    while candidato.lower() in usados:
        sufixo = f' ({n})'
        candidato = cortar_nome(base, 31 - len(sufixo)) + sufixo
        n += 1
    usados.add(candidato.lower())
    return candidato


def ordenar(df):
    """Ordem alfabética que respeita acentos (Ângela antes de Bruno)."""
    return df.sort_values('Nome', key=lambda s: s.map(normalizar))


def nome_do_arquivo(periodos, demo=False, hoje=None):
    """
    Nome dos arquivos exportados pelo mês dos vencimentos: ["01/09/2026 a 30/09/2026"] → Fechamento_UnifyData_2026-09.
    Vencimentos em meses diferentes viram "2026-09_a_2026-10"; sem período, vale a data de hoje.
    Na demonstração o nome termina em _DEMO, para o arquivo não passar por um fechamento real.
    """
    meses = sorted(f'{a}-{m}' for _, m, a in re.findall(r'(\d{2})/(\d{2})/(\d{4})', ' '.join(periodos or [])))
    if not meses:
        rotulo = (hoje or date.today()).isoformat()
    else:
        rotulo = meses[0] if meses[0] == meses[-1] else f'{meses[0]}_a_{meses[-1]}'
    return f"Fechamento_UnifyData_{rotulo}{'_DEMO' if demo else ''}"


def gerar_excel(merged_df, origem_cols, demo=False):
    """
    Gera o Excel final com uma aba por empresa.
    Colunas por origem: Vencido [X] | Total [X] | ... | Total Desconto
    merged_df deve conter colunas Vencido__<origem> e Total__<origem>.
    Na demonstração, cada aba avisa que os dados são fictícios (embaixo da tabela e no cabeçalho da impressão).
    """
    wb = Workbook()
    wb.remove(wb.active)

    header_fill = PatternFill(start_color='3FCFB9', end_color='3FCFB9', fill_type='solid')  # a cor principal do site, como no relatório
    header_font = Font(color='141416', bold=True)
    gray_font   = Font(color='808080')
    center      = Alignment(horizontal='center', vertical='center')
    moeda       = '"R$" #,##0.00'
    usados      = set()

    for emp in sorted(merged_df['Empresa'].dropna().unique(), key=normalizar):
        ws = wb.create_sheet(title=nome_de_aba(emp, usados))
        ws.freeze_panes = 'A2'
        # Impressão: A4, na largura da página, cabeçalho repetido; paisagem com 3 lojas ou mais, como o relatório
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = 'landscape' if len(origem_cols) >= 3 else 'portrait'
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = '1:1'

        # Cabeçalho: Funcionário | CPF | [Vencido X | Total X] | Total Desconto
        headers = ['Funcionário', 'CPF']
        for o in origem_cols:
            headers.append(f'Vencido {o}')
            headers.append(f'Total {o}')
        headers.append('Total Desconto')
        ws.append(headers)
        ws.row_dimensions[1].height = 20
        for cell in ws[1]:
            cell.fill = header_fill; cell.font = header_font; cell.alignment = center

        emp_df = ordenar(merged_df[merged_df['Empresa'] == emp])

        for _, row in emp_df.iterrows():
            row_data  = [row['Nome'], formatar_cpf(row['CPF'])]
            for o in origem_cols:
                row_data.append(float(row.get(f'Vencido__{o}', 0) or 0))
                row_data.append(float(row.get(f'Total__{o}',   0) or 0))
            row_data.append(float(row['Total Desconto']))
            ws.append(row_data)

            r = ws.max_row
            ws.cell(r, 1).data_type = 's'  # nome que começa com "=" fica texto, não vira fórmula
            for col_idx in range(3, len(row_data) + 1):
                ws.cell(r, col_idx).number_format = moeda
            if row['Total Desconto'] == 0:
                for col_idx in range(1, len(row_data) + 1):
                    ws.cell(r, col_idx).font = gray_font

        # Rodapé com a soma de cada coluna de valores
        somas = [float(emp_df[c].sum()) for o in origem_cols for c in (f'Vencido__{o}', f'Total__{o}')]
        ws.append(['TOTAL GERAL', ''] + [round(v, 2) for v in somas] + [round(float(emp_df['Total Desconto'].sum()), 2)])
        r = ws.max_row
        for col_idx in range(3, len(headers) + 1):
            ws.cell(r, col_idx).number_format = moeda
        for col_idx in range(1, len(headers) + 1):
            ws.cell(r, col_idx).font = Font(bold=True)
            ws.cell(r, col_idx).border = Border(top=Side(style='thin'))

        # Largura pelas maiores entradas, com teto para nomes muito longos.
        # Texto em maiúsculas ocupa mais que a unidade de largura do Excel (o "0"), daí a folga de 20%.
        for col in ws.columns:
            letra = col[0].column_letter
            tamanhos = [len(brl(c.value)) if isinstance(c.value, (int, float)) else math.ceil(len(str(c.value)) * 1.2)
                        for c in col if c.value is not None]
            ws.column_dimensions[letra].width = min(max(tamanhos, default=10) + 3, 48)

        # Depois das larguras, para o aviso longo não alargar a coluna dos nomes
        if demo:
            ws.append([])
            ws.append([AVISO_DEMO])
            ws.cell(ws.max_row, 1).font = Font(italic=True, color='808080')
            ws.oddHeader.center.text = 'DEMONSTRAÇÃO: dados fictícios'
            ws.oddHeader.center.font = '-,Bold'  # fonte padrão, em negrito

    if demo:
        wb.properties.subject = AVISO_DEMO
    if not wb.sheetnames:
        wb.create_sheet('Dados')

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def gerar_csv(merged_df, origem_cols):
    """CSV separado por ; e com vírgula decimal, como o Excel brasileiro espera."""
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=';', lineterminator='\r\n')
    cab = ['Empresa', 'Funcionário', 'CPF']
    for o in origem_cols:
        cab += [f'Vencido {o}', f'Total {o}']
    w.writerow([texto_seguro(c) for c in cab] + ['Total Desconto'])
    num = lambda v: f"{float(v or 0):.2f}".replace('.', ',')
    for _, row in merged_df.sort_values(['Empresa', 'Nome'], key=lambda s: s.map(normalizar)).iterrows():
        linha = [texto_seguro(row['Empresa']), texto_seguro(row['Nome']), formatar_cpf(row['CPF'])]
        for o in origem_cols:
            linha += [num(row.get(f'Vencido__{o}', 0)), num(row.get(f'Total__{o}', 0))]
        w.writerow(linha + [num(row['Total Desconto'])])
    return (BOM + buf.getvalue()).encode('utf-8')
