import streamlit as st
import pandas as pd
from thefuzz import process, fuzz
import io
import re
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment

# ──────────────────────────────────────────────
#  Configuração da Página
# ──────────────────────────────────────────────
st.set_page_config(page_title="UnifyData", page_icon="🔗", layout="wide")

st.markdown("""
<style>
    .stApp { background-color: #121214; color: #E0E0E0; }
    header { visibility: hidden; }
    .title-header {
        text-align: center; text-transform: uppercase; color: white;
        font-size: 2.5rem; font-weight: bold;
        margin-top: 1rem; margin-bottom: 0.5rem;
        font-family: 'Inter', sans-serif;
    }
    .title-line {
        border: none; height: 2px;
        background-color: #1E3A8A; width: 100%; margin-bottom: 3rem;
    }
    .footer {
        text-align: center; color: #666; font-size: 0.8rem;
        padding: 2rem 0; margin-top: 4rem;
        border-top: 1px solid #333; font-family: 'Inter', sans-serif;
    }
    .stButton > button {
        background-color: #1E3A8A !important; color: white !important;
        border-radius: 8px !important; border: none !important;
        padding: .5rem 2rem !important; font-weight: bold !important;
    }
    .stButton > button:hover { background-color: #254aab !important; }
</style>
""", unsafe_allow_html=True)

st.markdown('<div class="title-header">UnifyData</div>', unsafe_allow_html=True)
st.markdown('<div class="title-line"></div>', unsafe_allow_html=True)

# ──────────────────────────────────────────────
#  Funções Auxiliares
# ──────────────────────────────────────────────
def find_column(df, targets):
    """Localiza uma coluna pelo nome (cabeçalho ou primeiros valores)."""
    targets_lower = [t.strip().lower() for t in targets]
    for col in df.columns:
        if str(col).strip().lower() in targets_lower:
            return col
    # Busca nos primeiros 20 valores de cada coluna
    for col in df.columns:
        for val in df[col].head(20).values:
            if pd.notna(val) and str(val).strip().lower() in targets_lower:
                return col
    return None


def parse_hr_files(hr_files):
    """Lê arquivos de RH (XLS, XLSX ou CSV) e retorna DataFrame com Nome, CPF e Empresa."""
    records = []
    for file in hr_files:
        try:
            base = file.name.rsplit('.', 1)[0]
            if base.lower().startswith('empregados '):
                base = base[11:].strip()

            ext = file.name.rsplit('.', 1)[-1].lower()

            if ext == 'csv':
                # Tenta cp1252 primeiro (padrão Linx/Windows), depois utf-8
                try:
                    df = pd.read_csv(file, sep=';', encoding='cp1252', on_bad_lines='skip')
                except Exception:
                    file.seek(0)
                    df = pd.read_csv(file, sep=';', encoding='utf-8', on_bad_lines='skip')

                nome_col = find_column(df, ['nome', 'nome completo', 'funcionario', 'funcionário'])
                cpf_col  = find_column(df, ['cpf'])
                if not nome_col:
                    continue
                for _, row in df.iterrows():
                    nome = str(row[nome_col]).strip().upper()
                    if not nome or nome in ('NAN', 'NOME', 'NOME COMPLETO', 'FUNCIONARIO', 'FUNCIONÁRIO') or len(nome) < 3:
                        continue
                    cpf = ''
                    if cpf_col:
                        cpf = re.sub(r'[^\d]', '', str(row[cpf_col]))
                    records.append({'Nome': ' '.join(nome.split()), 'CPF': cpf, 'Empresa': base})

            else:
                xls = pd.ExcelFile(file)
                for sheet in xls.sheet_names:
                    empresa = f"{base} - {sheet}" if (len(xls.sheet_names) > 1 and not sheet.lower().startswith('plan')) else base
                    df = pd.read_excel(xls, sheet_name=sheet)
                    nome_col = find_column(df, ['nome', 'nome completo', 'funcionario', 'funcionário'])
                    cpf_col  = find_column(df, ['cpf'])
                    if not nome_col:
                        continue
                    for _, row in df.iterrows():
                        nome = str(row[nome_col]).strip().upper()
                        if not nome or nome in ('NAN', 'NOME', 'NOME COMPLETO', 'FUNCIONARIO', 'FUNCIONÁRIO') or len(nome) < 3:
                            continue
                        cpf = ''
                        if cpf_col:
                            cpf = re.sub(r'[^\d]', '', str(row[cpf_col]))
                        records.append({'Nome': ' '.join(nome.split()), 'CPF': cpf, 'Empresa': empresa})

        except Exception as e:
            st.error(f"Erro ao processar {file.name}: {e}")

    return pd.DataFrame(records) if records else pd.DataFrame(columns=['Nome', 'CPF', 'Empresa'])


def parse_linx_files(linx_files):
    """
    Lê relatórios CSV do Linx (Pendências por Responsável).
    Retorna DataFrame com Nome, Vencido, Total e Origem —
    capturando os DOIS valores da linha Total separadamente.
    """
    def to_float(s):
        try:
            return float(re.sub(r'[R$\s]', '', s).replace('.', '').replace(',', '.'))
        except (ValueError, AttributeError):
            return 0.0

    records = []
    for file in linx_files:
        try:
            try:
                content = file.getvalue().decode('cp1252')
            except Exception:
                content = file.getvalue().decode('utf-8', errors='ignore')

            origem       = re.sub(r'\.(csv|txt)$', '', file.name, flags=re.IGNORECASE)
            current_nome = None

            for line in content.split('\n'):
                if 'Respons' in line:
                    m = re.search(r'Respons[^\s:]*:\s*\d+\s*-\s*([^;\r\n]+)', line, re.IGNORECASE)
                    if m:
                        current_nome = ' '.join(m.group(1).strip().upper().split())

                if current_nome and line.strip().lower().startswith('total'):
                    if re.match(r'total\s+geral', line.strip(), re.IGNORECASE):
                        current_nome = None
                        continue
                    parts = [p.strip() for p in line.split(';') if p.strip()]
                    # parts[0]="Total"  parts[1]=vencido  parts[2]=total geral
                    vencido = to_float(parts[1]) if len(parts) >= 2 else 0.0
                    total   = to_float(parts[2]) if len(parts) >= 3 else vencido
                    if total > 0:
                        records.append({
                            'Nome': current_nome,
                            'Vencido': vencido,
                            'Total': total,
                            'Origem': origem
                        })
                    current_nome = None

        except Exception as e:
            st.error(f"Erro ao processar {file.name}: {e}")

    return pd.DataFrame(records) if records else pd.DataFrame(columns=['Nome', 'Vencido', 'Total', 'Origem'])


def gerar_excel(merged_df, origem_cols):
    """
    Gera o Excel final com uma aba por empresa.
    Colunas por origem: Vencido [X] | Total [X] | ... | Total Desconto
    merged_df deve conter colunas Vencido__<origem> e Total__<origem>.
    """
    wb = Workbook()
    wb.remove(wb.active)

    header_fill = PatternFill(start_color='58111A', end_color='58111A', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)
    gray_font   = Font(color='808080')
    center      = Alignment(horizontal='center', vertical='center')

    for emp in sorted(merged_df['Empresa'].dropna().unique()):
        ws = wb.create_sheet(title=str(emp)[:31])

        # Cabeçalho: Funcionário | CPF | [Vencido X | Total X] | Total Desconto
        headers = ['Funcionário', 'CPF']
        for o in origem_cols:
            headers.append(f'Vencido {o}')
            headers.append(f'Total {o}')
        headers.append('Total Desconto')
        ws.append(headers)
        for cell in ws[1]:
            cell.fill = header_fill; cell.font = header_font; cell.alignment = center

        emp_df    = merged_df[merged_df['Empresa'] == emp].sort_values('Nome')
        grand_sum = 0

        for _, row in emp_df.iterrows():
            row_data  = [row['Nome'], row['CPF']]
            row_total = 0
            for o in origem_cols:
                v = row.get(f'Vencido__{o}', 0) or 0
                t = row.get(f'Total__{o}',   0) or 0
                row_data.append(v)
                row_data.append(t)
                row_total += t
            row_data.append(row_total)
            grand_sum += row_total
            ws.append(row_data)

            r = ws.max_row
            for col_idx in range(3, len(row_data) + 1):
                ws.cell(r, col_idx).number_format = 'R$ #,##0.00'
            if row_total == 0:
                for col_idx in range(1, len(row_data) + 1):
                    ws.cell(r, col_idx).font = gray_font

        # Rodapé
        footer = ['TOTAL GERAL', ''] + ['' for _ in range(len(origem_cols) * 2)] + [grand_sum]
        ws.append(footer)
        r = ws.max_row
        ws.cell(r, 1).font = Font(bold=True)
        ws.cell(r, len(footer)).font = Font(bold=True)
        ws.cell(r, len(footer)).number_format = 'R$ #,##0.00'

        # Auto-largura
        for col in ws.columns:
            col_letter = col[0].column_letter
            max_len = max((len(str(c.value)) for c in col if c.value), default=10)
            ws.column_dimensions[col_letter].width = max_len + 3

    if not wb.sheetnames:
        wb.create_sheet('Dados')

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ──────────────────────────────────────────────
#  Estado da Sessão
# ──────────────────────────────────────────────
defaults = {
    'processing_done': False,
    'hr_df': None,
    'linx_df': None,
    'linx_vencido': None,
    'linx_total': None,
    'divergences': [],
    'origens_upload': [],
}
for k, v in defaults.items():
    if k not in st.session_state:
        st.session_state[k] = v


# ──────────────────────────────────────────────
#  Interface: Upload
# ──────────────────────────────────────────────
if not st.session_state.processing_done:
    col1, col2 = st.columns(2)

    with col1:
        st.markdown('### 📋 Arquivos do RH')
        st.caption('Planilhas de funcionários (Excel ou CSV)')
        hr_files = st.file_uploader(
            'Arraste os arquivos de RH', accept_multiple_files=True,
            type=['xlsx', 'xls', 'csv'], label_visibility='collapsed'
        )

    with col2:
        st.markdown('### 🛒 Relatórios Linx')
        st.caption('Pendências por Responsável extraídas do Linx (.csv)')
        linx_files = st.file_uploader(
            'Arraste os relatórios do Linx', accept_multiple_files=True,
            type=['csv', 'txt'], label_visibility='collapsed'
        )

    st.markdown('<br>', unsafe_allow_html=True)

    if st.button('⚙️  Processar e Unificar Dados', use_container_width=True):
        if not hr_files or not linx_files:
            st.warning('⚠️ Envie os arquivos de RH **e** do Linx para continuar.')
        else:
            with st.spinner('Analisando e cruzando informações...'):
                origens = [f.name.rsplit('.', 1)[0] for f in linx_files]
                st.session_state.origens_upload = origens

                hr_df   = parse_hr_files(hr_files)
                linx_df = parse_linx_files(linx_files)

                if hr_df.empty:
                    st.error('Não foi possível extrair dados dos arquivos de RH. Verifique se há colunas "Nome" e "CPF".')
                elif linx_df.empty:
                    st.error('Não foi possível extrair dados do Linx. Verifique o layout dos arquivos.')
                else:
                    # Agrega Linx: pivot separado para Vencido e Total por origem
                    vencido_agg = linx_df.groupby(['Nome', 'Origem'], as_index=False)['Vencido'].sum()
                    total_agg   = linx_df.groupby(['Nome', 'Origem'], as_index=False)['Total'].sum()

                    hr_names = hr_df['Nome'].tolist()

                    divergences = []
                    for nome in linx_df['Nome'].unique():
                        if nome in hr_names:
                            continue
                        best, score = process.extractOne(nome, hr_names, scorer=fuzz.token_sort_ratio)
                        total_val   = total_agg[total_agg['Nome'] == nome]['Total'].sum()
                        origens_str = ', '.join(linx_df[linx_df['Nome'] == nome]['Origem'].unique())
                        divergences.append({
                            'Linx_Nome': nome, 'HR_Nome': best, 'Score': score,
                            'Tipo': 'Nome Semelhante' if score >= 85 else 'Não Encontrado',
                            'Origem': origens_str, 'Valor': total_val
                        })

                    st.session_state.hr_df        = hr_df
                    st.session_state.linx_vencido = vencido_agg
                    st.session_state.linx_total   = total_agg
                    st.session_state.linx_df      = linx_df   # raw (para filtros nas divergências)
                    st.session_state.divergences  = divergences
                    st.session_state.processing_done = True
                    st.rerun()


# ──────────────────────────────────────────────
#  Interface: Divergências
# ──────────────────────────────────────────────
if st.session_state.processing_done:
    divs = st.session_state.divergences

    if divs:
        st.markdown('### ⚠️ Central de Revisão e Divergências')
        st.info(f'{len(divs)} divergência(s) encontrada(s). Resolva-as antes de gerar o relatório.')

        # Funções definidas UMA VEZ fora do loop — evita bug de closure do Streamlit
        def renomear(nome_antigo, nome_novo):
            for key in ('linx_df', 'linx_vencido', 'linx_total'):
                df = st.session_state[key]
                if df is not None:
                    df.loc[df['Nome'] == nome_antigo, 'Nome'] = nome_novo
                    st.session_state[key] = df

        def remover(nome):
            for key in ('linx_df', 'linx_vencido', 'linx_total'):
                df = st.session_state[key]
                if df is not None:
                    st.session_state[key] = df[df['Nome'] != nome]

        for i, div in enumerate(divs):

                if div['Tipo'] == 'Nome Semelhante':
                    st.warning(
                        f"**Nome semelhante detectado** → Linx: **{div['Linx_Nome']}** "
                        f"/ RH: **{div['HR_Nome']}** (similaridade: {div['Score']}%)"
                    )
                    c1, c2, c3 = st.columns(3)
                    with c1:
                        if st.button('✅ Aprovar vínculo', key=f'apv_{i}'):
                            renomear(div['Linx_Nome'], div['HR_Nome'])
                            st.session_state.divergences.pop(i)
                            st.rerun()
                    with c2:
                        if st.button('🚫 Ignorar', key=f'ign_{i}'):
                            remover(div['Linx_Nome'])
                            st.session_state.divergences.pop(i)
                            st.rerun()
                    with c3:
                        cpf_in = st.text_input('Vincular por CPF', key=f'cpf_{i}', placeholder='000.000.000-00')
                        if st.button('🔗 Confirmar CPF', key=f'cpf_btn_{i}'):
                            match = st.session_state.hr_df[st.session_state.hr_df['CPF'] == re.sub(r'[^\d]', '', cpf_in)]
                            if not match.empty:
                                renomear(div['Linx_Nome'], match.iloc[0]['Nome'])
                                st.session_state.divergences.pop(i)
                                st.rerun()
                            else:
                                st.error('CPF não encontrado no RH.')

                elif div['Tipo'] == 'Não Encontrado':
                    st.error(f"**Não encontrado no RH:** **{div['Linx_Nome']}** não existe em nenhuma lista.")
                    c1, c2 = st.columns([3, 1])
                    with c1:
                        empresas = sorted(st.session_state.hr_df['Empresa'].unique())
                        cpf_nf   = st.text_input('CPF (obrigatório)', key=f'cpf_nf_{i}', placeholder='000.000.000-00')
                        emp_nf   = st.selectbox('Empresa', empresas, key=f'emp_nf_{i}')
                    with c2:
                        st.markdown('<br><br>', unsafe_allow_html=True)
                        if st.button('🚫 Ignorar', key=f'ign_nf_{i}'):
                            remover(div['Linx_Nome'])
                            st.session_state.divergences.pop(i)
                            st.rerun()

                    if st.button('💾 Criar registro e aprovar', key=f'create_nf_{i}'):
                        cpf_clean = re.sub(r'[^\d]', '', cpf_nf)
                        if not cpf_clean:
                            st.error('Preencha o CPF antes de registrar.')
                        else:
                            new_row = pd.DataFrame([{'Nome': div['Linx_Nome'], 'CPF': cpf_clean, 'Empresa': emp_nf}])
                            st.session_state.hr_df = pd.concat([st.session_state.hr_df, new_row], ignore_index=True)
                            st.session_state.divergences.pop(i)
                            st.rerun()

    # ──────────────────────────────────────────────
    #  Interface: Download
    # ──────────────────────────────────────────────
    else:
        st.success('✅ Tudo certo! Nenhuma divergência pendente.')

        hr      = st.session_state.hr_df
        origens = st.session_state.origens_upload
        v_agg   = st.session_state.linx_vencido  # cols: Nome, Origem, Vencido
        t_agg   = st.session_state.linx_total    # cols: Nome, Origem, Total

        # Pivot separado para Vencido e Total
        def make_pivot(df, col):
            if df is None or df.empty:
                return pd.DataFrame(columns=['Nome'])
            p = df.pivot_table(index='Nome', columns='Origem', values=col, aggfunc='sum', fill_value=0)
            # Renomeia as colunas ANTES do reset_index para evitar conflito com 'Nome'
            p.columns = [f'{col}__{o}' for o in p.columns]
            return p.reset_index()

        piv_v = make_pivot(v_agg, 'Vencido')
        piv_t = make_pivot(t_agg, 'Total')

        merged = pd.merge(hr, piv_v, on='Nome', how='left')
        merged = pd.merge(merged, piv_t, on='Nome', how='left')
        merged = merged.fillna(0)

        buf = gerar_excel(merged, origens)

        st.markdown('<br>', unsafe_allow_html=True)
        st.download_button(
            label='⬇️  Baixar Relatório Consolidado (.xlsx)',
            data=buf,
            file_name='Relatorio_UnifyData.xlsx',
            mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            use_container_width=True
        )

        st.markdown('<br>', unsafe_allow_html=True)
        if st.button('🔄 Reiniciar Sistema', use_container_width=True):
            for k in list(defaults.keys()):
                st.session_state.pop(k, None)
            st.rerun()

st.markdown('<div class="footer">POWERED BY AFN SYSTEMS</div>', unsafe_allow_html=True)
